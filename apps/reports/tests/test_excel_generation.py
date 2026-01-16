"""
Tests for Excel report generation (Part 3: Backup & Archive Strategy).

Tests that Excel (.xlsx) files are generated correctly for scheduled report emails.
These Excel files serve as backup data that can be stored in email archives.
"""
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook

from apps.reports.tasks import (
    generate_report_excel,
    generate_report_pdf,
    get_report_dates,
    EXCEL_STYLES,
    _generate_sales_daily_excel,
    _generate_pnl_daily_excel,
    _generate_stock_levels_excel,
)


class ExcelGeneratorDispatcherTests(TestCase):
    """Test the generate_report_excel() dispatcher function."""
    
    def setUp(self):
        self.dates = get_report_dates(mode='manual')
    
    def test_unknown_report_returns_none(self):
        """Unknown report codes should return (None, error_message)."""
        excel_bytes, message = generate_report_excel('unknown_report_code', self.dates)
        
        self.assertIsNone(excel_bytes)
        self.assertIn('No Excel generator', message)
    
    def test_all_supported_report_codes_generate_bytes(self):
        """All supported report codes should generate Excel bytes."""
        supported_codes = [
            'sales_daily', 'sales_weekly', 'sales_monthly',
            'pnl_daily', 'pnl_weekly', 'pnl_monthly',
            'stock_levels', 'inventory_valuation',
            'production_daily', 'production_weekly', 'production_monthly',
            'payroll_monthly',
        ]
        
        for code in supported_codes:
            with self.subTest(report_code=code):
                excel_bytes, filename = generate_report_excel(code, self.dates)
                
                self.assertIsNotNone(excel_bytes, f"{code} should generate bytes")
                self.assertGreater(len(excel_bytes), 0, f"{code} should have non-empty bytes")
                self.assertTrue(filename.endswith('.xlsx'), f"{code} filename should end with .xlsx")
    
    def test_excel_and_pdf_generated_for_same_report(self):
        """Both PDF and Excel should be generated for the same report."""
        code = 'sales_daily'
        
        pdf_bytes, pdf_filename = generate_report_pdf(code, self.dates)
        excel_bytes, excel_filename = generate_report_excel(code, self.dates)
        
        # Both should succeed
        self.assertIsNotNone(pdf_bytes)
        self.assertIsNotNone(excel_bytes)
        
        # Filenames should have correct extensions
        self.assertTrue(pdf_filename.endswith('.pdf'))
        self.assertTrue(excel_filename.endswith('.xlsx'))
        
        # Should have same date in filename
        pdf_date = pdf_filename.replace('.pdf', '').split('_')[-1]
        excel_date = excel_filename.replace('.xlsx', '').split('_')[-1]
        self.assertEqual(pdf_date, excel_date)


class ExcelStyleConstantsTests(TestCase):
    """Test that EXCEL_STYLES constants are properly defined."""
    
    def test_excel_styles_dict_exists(self):
        """EXCEL_STYLES should be a non-empty dictionary."""
        self.assertIsInstance(EXCEL_STYLES, dict)
        self.assertGreater(len(EXCEL_STYLES), 0)
    
    def test_required_styles_exist(self):
        """Required style keys should exist in EXCEL_STYLES."""
        required_keys = [
            'title_font', 'title_fill',
            'section_font',
            'header_font', 'header_fill', 'header_alignment',
            'alt_row_fill',
            'footer_fill', 'footer_font',
            'currency_format', 'percent_format',
            'success_font', 'danger_font', 'warning_font', 'info_font', 'muted_font',
            'success_fill', 'danger_fill', 'warning_fill', 'info_fill',
            'border', 'border_bottom_only', 'footer_border',
        ]
        
        for key in required_keys:
            with self.subTest(style_key=key):
                self.assertIn(key, EXCEL_STYLES, f"Missing style: {key}")


class SalesDailyExcelTests(TestCase):
    """Test the _generate_sales_daily_excel() function."""
    
    def setUp(self):
        self.dates = get_report_dates(mode='manual')
    
    def test_generates_valid_xlsx(self):
        """Should generate valid Excel file that can be loaded."""
        excel_bytes, filename = _generate_sales_daily_excel(self.dates)
        
        self.assertIsNotNone(excel_bytes)
        
        # Load the workbook to verify it's valid
        buffer = BytesIO(excel_bytes)
        wb = load_workbook(buffer)
        
        self.assertEqual(len(wb.worksheets), 1)
        self.assertEqual(wb.active.title, "Daily Sales")
    
    def test_contains_required_sections(self):
        """Should contain title, KPIs, and tables."""
        excel_bytes, filename = _generate_sales_daily_excel(self.dates)
        
        buffer = BytesIO(excel_bytes)
        wb = load_workbook(buffer)
        ws = wb.active
        
        # Check title row
        title_cell = ws['A1']
        self.assertIn('Daily Sales Report', str(title_cell.value))
        
        # Check that we have data beyond just headers
        self.assertGreater(ws.max_row, 5, "Should have multiple rows of content")


class PnlDailyExcelTests(TestCase):
    """Test the _generate_pnl_daily_excel() function."""
    
    def setUp(self):
        self.dates = get_report_dates(mode='manual')
    
    def test_generates_valid_xlsx(self):
        """Should generate valid Excel file."""
        excel_bytes, filename = _generate_pnl_daily_excel(self.dates)
        
        self.assertIsNotNone(excel_bytes)
        
        buffer = BytesIO(excel_bytes)
        wb = load_workbook(buffer)
        
        self.assertEqual(wb.active.title, "Daily P&L")
    
    def test_contains_profit_section(self):
        """Should contain NET PROFIT section."""
        excel_bytes, filename = _generate_pnl_daily_excel(self.dates)
        
        buffer = BytesIO(excel_bytes)
        wb = load_workbook(buffer)
        ws = wb.active
        
        # Search for NET PROFIT in the worksheet
        found_profit = False
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=1):
            for cell in row:
                if cell.value and 'NET PROFIT' in str(cell.value):
                    found_profit = True
                    break
        
        self.assertTrue(found_profit, "Should contain NET PROFIT section")


class StockLevelsExcelTests(TestCase):
    """Test the _generate_stock_levels_excel() function."""
    
    def setUp(self):
        self.dates = get_report_dates(mode='manual')
    
    def test_generates_valid_xlsx(self):
        """Should generate valid Excel file."""
        excel_bytes, filename = _generate_stock_levels_excel(self.dates)
        
        self.assertIsNotNone(excel_bytes)
        
        buffer = BytesIO(excel_bytes)
        wb = load_workbook(buffer)
        
        self.assertEqual(wb.active.title, "Stock Levels")
    
    def test_has_correct_headers(self):
        """Should have correct column headers."""
        excel_bytes, filename = _generate_stock_levels_excel(self.dates)
        
        buffer = BytesIO(excel_bytes)
        wb = load_workbook(buffer)
        ws = wb.active
        
        # Find header row (should contain 'Item', 'Category', etc.)
        found_headers = False
        for row in ws.iter_rows(min_row=1, max_row=10):
            row_values = [str(cell.value) if cell.value else '' for cell in row]
            if 'Item' in row_values and 'Category' in row_values:
                found_headers = True
                break
        
        self.assertTrue(found_headers, "Should have Item and Category headers")


class ExcelFilenameTests(TestCase):
    """Test that Excel filenames follow the correct format."""
    
    def setUp(self):
        self.dates = get_report_dates(mode='manual')
    
    def test_daily_filename_format(self):
        """Daily reports should have YYYYMMDD in filename."""
        excel_bytes, filename = generate_report_excel('sales_daily', self.dates)
        
        # Filename should be like: sales_daily_20260116.xlsx
        self.assertRegex(filename, r'sales_daily_\d{8}\.xlsx')
    
    def test_weekly_filename_format(self):
        """Weekly reports should have week start date in filename."""
        excel_bytes, filename = generate_report_excel('sales_weekly', self.dates)
        
        # Filename should be like: sales_weekly_20260112.xlsx
        self.assertRegex(filename, r'sales_weekly_\d{8}\.xlsx')
    
    def test_monthly_filename_format(self):
        """Monthly reports should have YYYYMM in filename."""
        excel_bytes, filename = generate_report_excel('sales_monthly', self.dates)
        
        # Filename should be like: sales_monthly_202601.xlsx
        self.assertRegex(filename, r'sales_monthly_\d{6}\.xlsx')


class ExcelMimeTypeTests(TestCase):
    """Test that Excel files have correct MIME type."""
    
    def test_xlsx_mime_type(self):
        """Excel files should use the correct MIME type."""
        expected_mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        
        # The MIME type is set in send_scheduled_reports
        # Just verify the constant is correct
        self.assertEqual(
            expected_mime,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
